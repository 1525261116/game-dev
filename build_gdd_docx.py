# -*- coding: utf-8 -*-
"""把《江湖从心》GDD 各章 Markdown 合并为单一交付稿，并导出 DOCX。

用法：
    python3 build_gdd_docx.py

产出（均位于项目根 /Users/zq/game-dev/）：
    江湖从心_完整策划设计_v1.0.md / .docx
    江湖从心_引擎选型评估报告_v1.0.md / .docx
"""

from __future__ import annotations

import re
import sys
from pathlib import Path

from docx import Document
from docx.enum.table import WD_TABLE_ALIGNMENT
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.oxml.ns import qn
from docx.shared import Pt, RGBColor, Cm

ROOT = Path(__file__).resolve().parent
GDD = ROOT / "jianghu-design" / "GDD"
TECH = ROOT / "jianghu-design" / "12_tech"

CJK_FONT = "PingFang SC"
MONO_FONT = "Menlo"
INK = RGBColor(0x1F, 0x1F, 0x1F)
ACCENT = RGBColor(0x8C, 0x6B, 0x2F)   # 暗金
RED = RGBColor(0xA6, 0x2B, 0x1F)      # 朱红

GDD_CHAPTERS = [
    ("00_总览与设计目标.md", "第 0 章 总览与设计目标"),
    ("01_玩法与核心循环.md", "第 1 章 玩法与核心循环"),
    ("02_系统设计.md", "第 2 章 系统设计"),
    ("03_数值框架.md", "第 3 章 数值框架"),
    ("04_关卡与内容规划.md", "第 4 章 关卡与内容规划"),
    ("05_生产计划与执行.md", "第 5 章 生产计划与执行"),
]


# --------------------------------------------------------------------- 合并 MD
def build_merged_md() -> Path:
    parts = [
        "# 江湖从心 · 完整策划设计 v1.0（执行版）",
        "",
        "> 武侠开放世界 · 单机 RPG · 多端（移动端 + PC）",
        "> 依据《江湖从心_策划案_v0.2》与 `jianghu-design/` 全部设计稿收敛而成",
        "> 配套文档：《江湖从心_引擎选型评估报告_v1.0》",
        "",
        "---",
        "",
        "## 目录",
        "",
    ]
    for name, title in GDD_CHAPTERS:
        parts.append(f"- {title}")
    parts += [
        "",
        "## 如何使用本文件",
        "",
        "| 章节 | 内容 | 主要读者 |",
        "|------|------|----------|",
        "| 第 0 章 | 产品定位、商业化、平台策略、冲突裁决、术语表 | 制作人 / 主策 |",
        "| 第 1 章 | 玩法模块、四层核心循环、首小时体验、多周目 | 全体 |",
        "| 第 2 章 | 14 个系统的规则 / 数值 / 界面 / 验收标准 | 系统策划 |",
        "| 第 3 章 | 战斗公式、成长曲线、经济平衡、时间系统、验收 | 数值策划 |",
        "| 第 4 章 | 区域与关卡规划、内容量、配置表契约、工时估算 | 关卡策划 / 主程 |",
        "| 第 5 章 | 里程碑、团队、技术预算、回本测算、风险登记册 | 制作人 |",
        "",
        "> 本文件与既有设计稿冲突时，**以本文件为准**（冲突裁决表见第 0 章 0.9 节）。",
        "",
        "---",
        "",
    ]
    for name, _ in GDD_CHAPTERS:
        text = (GDD / name).read_text(encoding="utf-8").strip()
        parts.append(text)
        parts += ["", "---", ""]
    out = ROOT / "江湖从心_完整策划设计_v1.0.md"
    out.write_text("\n".join(parts).rstrip() + "\n", encoding="utf-8")
    return out


# ------------------------------------------------------------------ MD → DOCX
INLINE_RE = re.compile(
    r"(\*\*.+?\*\*|`[^`]+`|\[[^\]]+\]\([^)]+\)|~~.+?~~)"
)
LINK_RE = re.compile(r"\[([^\]]+)\]\(([^)]+)\)")


def add_runs(paragraph, text: str, base_size: float = 10.5, mono: bool = False) -> None:
    """按内联标记写入 run（加粗 / 行内代码 / 链接 / 删除线）。"""
    for token in INLINE_RE.split(text):
        if not token:
            continue
        bold = False
        italic = False
        strike = False
        mono_run = mono
        content = token
        if token.startswith("**") and token.endswith("**") and len(token) > 4:
            bold, content = True, token[2:-2]
        elif token.startswith("~~") and token.endswith("~~") and len(token) > 4:
            strike, content = True, token[2:-2]
        elif token.startswith("`") and token.endswith("`") and len(token) > 2:
            mono_run, content = True, token[1:-1]
        else:
            m = LINK_RE.fullmatch(token)
            if m:
                label, url = m.group(1), m.group(2)
                content = f"{label} ({url})" if url.startswith("http") else label
        run = paragraph.add_run(content)
        run.bold = bold
        run.italic = italic
        run.font.strike = strike
        run.font.size = Pt(base_size if not mono_run else base_size - 0.5)
        run.font.color.rgb = INK
        name = MONO_FONT if mono_run else CJK_FONT
        run.font.name = name
        run._element.rPr.rFonts.set(qn("w:eastAsia"), name)
        if mono_run:
            run.font.color.rgb = RGBColor(0x33, 0x44, 0x55)


def style_document(document: Document) -> None:
    normal = document.styles["Normal"]
    normal.font.name = CJK_FONT
    normal.font.size = Pt(10.5)
    normal.font.color.rgb = INK
    normal._element.rPr.rFonts.set(qn("w:eastAsia"), CJK_FONT)
    normal.paragraph_format.space_after = Pt(6)
    normal.paragraph_format.line_spacing = 1.25

    for name, size, color, bold in (
        ("Heading 1", 18, ACCENT, True),
        ("Heading 2", 15, ACCENT, True),
        ("Heading 3", 12.5, INK, True),
        ("Heading 4", 11, INK, True),
    ):
        st = document.styles[name]
        st.font.name = CJK_FONT
        st.font.size = Pt(size)
        st.font.bold = bold
        st.font.color.rgb = color
        st._element.rPr.rFonts.set(qn("w:eastAsia"), CJK_FONT)
        st.paragraph_format.space_before = Pt(12 if name != "Heading 1" else 18)
        st.paragraph_format.space_after = Pt(6)

    for sec in document.sections:
        sec.top_margin = Cm(2.0)
        sec.bottom_margin = Cm(2.0)
        sec.left_margin = Cm(2.2)
        sec.right_margin = Cm(2.2)


def add_table(document: Document, rows: list[list[str]]) -> None:
    if not rows:
        return
    cols = max(len(r) for r in rows)
    table = document.add_table(rows=0, cols=cols)
    table.style = "Table Grid"
    table.alignment = WD_TABLE_ALIGNMENT.CENTER
    table.autofit = True
    for i, row in enumerate(rows):
        cells = table.add_row().cells
        for j in range(cols):
            text = row[j] if j < len(row) else ""
            cell = cells[j]
            cell.text = ""
            p = cell.paragraphs[0]
            p.paragraph_format.space_after = Pt(2)
            p.paragraph_format.line_spacing = 1.05
            add_runs(p, text, base_size=8.5)
            if i == 0:
                for run in p.runs:
                    run.bold = True
                    run.font.color.rgb = RGBColor(0xFF, 0xFF, 0xFF)
                tcPr = cell._tc.get_or_add_tcPr()
                shd = tcPr.makeelement(qn("w:shd"), {
                    qn("w:val"): "clear", qn("w:color"): "auto", qn("w:fill"): "6B5B3E"})
                tcPr.append(shd)
    document.add_paragraph()


def md_to_docx(md_path: Path, docx_path: Path, title: str, subtitle: str) -> None:
    document = Document()
    style_document(document)

    document.add_heading(title, level=0)
    for line in subtitle.splitlines():
        p = document.add_paragraph()
        p.alignment = WD_ALIGN_PARAGRAPH.LEFT
        add_runs(p, line, base_size=10.5)

    lines = md_path.read_text(encoding="utf-8").splitlines()
    i = 0
    in_code = False
    code_buf: list[str] = []
    first_h1_seen = False

    def flush_code() -> None:
        nonlocal code_buf
        if not code_buf:
            return
        for cl in code_buf:
            p = document.add_paragraph()
            p.paragraph_format.space_after = Pt(0)
            p.paragraph_format.line_spacing = 1.0
            p.paragraph_format.left_indent = Cm(0.6)
            run = p.add_run(cl if cl.strip() else " ")
            run.font.name = MONO_FONT
            run._element.rPr.rFonts.set(qn("w:eastAsia"), MONO_FONT)
            run.font.size = Pt(9)
            run.font.color.rgb = RGBColor(0x2A, 0x3A, 0x4A)
        code_buf = []
        document.add_paragraph()

    while i < len(lines):
        raw = lines[i]
        line = raw.rstrip()

        if line.strip().startswith("```"):
            if in_code:
                in_code = False
                flush_code()
            else:
                in_code = True
                code_buf = []
            i += 1
            continue
        if in_code:
            code_buf.append(raw)
            i += 1
            continue

        stripped = line.strip()

        # 表格
        if stripped.startswith("|") and stripped.count("|") >= 2:
            rows: list[list[str]] = []
            while i < len(lines):
                cur = lines[i].strip()
                if not (cur.startswith("|") and cur.count("|") >= 2):
                    break
                cells = [c.strip() for c in cur.strip("|").split("|")]
                if all(re.fullmatch(r":?-{2,}:?", c) or c == "" for c in cells):
                    i += 1
                    continue
                rows.append(cells)
                i += 1
            add_table(document, rows)
            continue

        # 分隔线
        if re.fullmatch(r"-{3,}", stripped):
            i += 1
            continue

        # 标题
        m = re.match(r"^(#{1,6})\s+(.*)$", stripped)
        if m:
            level = min(len(m.group(1)), 4)
            text = re.sub(r"\*\*", "", m.group(2))
            if level == 1 and not first_h1_seen:
                # 文档标题已由 add_heading 写入，跳过 Markdown 的首个 H1
                first_h1_seen = True
                i += 1
                continue
            document.add_heading("", level=level).add_run(text)
            i += 1
            continue

        # 引用
        if stripped.startswith(">"):
            buf = []
            while i < len(lines) and lines[i].strip().startswith(">"):
                buf.append(lines[i].strip()[1:].strip())
                i += 1
            p = document.add_paragraph()
            p.paragraph_format.left_indent = Cm(0.8)
            p.paragraph_format.space_before = Pt(4)
            add_runs(p, " ".join(buf), base_size=10)
            for run in p.runs:
                run.font.color.rgb = RGBColor(0x44, 0x44, 0x44)
            continue

        # 列表
        m = re.match(r"^(\s*)([-*+]|\d+\.)\s+(.*)$", line)
        if m:
            indent, _marker, body = m.groups()
            depth = len(indent) // 2
            style = "List Bullet" if _marker in "-*+" else "List Number"
            if depth >= 1:
                style = f"{style} {min(depth + 1, 3)}"
            try:
                p = document.add_paragraph(style=style)
            except KeyError:
                p = document.add_paragraph()
            p.paragraph_format.space_after = Pt(3)
            add_runs(p, body)
            i += 1
            continue

        # 空行
        if not stripped:
            i += 1
            continue

        # 普通段落
        p = document.add_paragraph()
        add_runs(p, stripped)
        i += 1

    document.save(docx_path)


def main() -> int:
    merged = build_merged_md()
    print(f"[ok] merged markdown -> {merged}")

    engine_md = TECH / "引擎选型评估报告_v1.0.md"
    engine_copy = ROOT / "江湖从心_引擎选型评估报告_v1.0.md"
    engine_copy.write_text(engine_md.read_text(encoding="utf-8"), encoding="utf-8")
    print(f"[ok] engine markdown -> {engine_copy}")

    md_to_docx(
        merged,
        ROOT / "江湖从心_完整策划设计_v1.0.docx",
        "江湖从心 · 完整策划设计 v1.0（执行版）",
        "武侠开放世界 · 单机 RPG · 多端（移动端 + PC）\n"
        "依据《江湖从心_策划案_v0.2》与 jianghu-design 全部设计稿收敛而成\n"
        "配套文档：江湖从心_引擎选型评估报告_v1.0",
    )
    print("[ok] -> 江湖从心_完整策划设计_v1.0.docx")

    md_to_docx(
        engine_copy,
        ROOT / "江湖从心_引擎选型评估报告_v1.0.docx",
        "江湖从心 · 引擎选型评估报告 v1.0",
        "面向移动端 + PC 双端的引擎选型、跨端发布能力与落地成本评估\n"
        "报告日期：2026-10-11｜首选建议：Unity 6.x LTS（备选 Cocos Creator）",
    )
    print("[ok] -> 江湖从心_引擎选型评估报告_v1.0.docx")
    return 0


if __name__ == "__main__":
    sys.exit(main())
